"""Heterogeneous document loading for the SupportPearlz knowledge base.

Recursively scans a knowledge-base directory, dispatches each file to a
format-specific loader, normalizes text, attaches consistent metadata, skips
empty/duplicate content, and tolerates individual corrupt files without
aborting the whole ingestion run.
"""

from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import docx  # python-docx
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document

from src.utils.logging_setup import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".csv"}

_DOC_TYPE_KEYWORDS = [
    ("warranty", "warranty_policy"),
    ("refund", "refund_returns_policy"),
    ("return", "refund_returns_policy"),
    ("shipping", "shipping_policy"),
    ("install", "installation_guide"),
    ("troubleshoot", "troubleshooting_guide"),
    ("faq", "faq"),
    ("pricing", "pricing_guide"),
    ("privacy", "privacy_policy"),
    ("service", "service_maintenance_agreement"),
    ("maintenance", "service_maintenance_agreement"),
    ("manual", "product_manual"),
]

_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.*)$")
_LAST_UPDATED_PATTERN = re.compile(r"last[\s_-]?updated\s*:\s*(.+)", re.IGNORECASE)
_WHITESPACE_PATTERN = re.compile(r"[ \t]+")
_BLANK_LINES_PATTERN = re.compile(r"\n{3,}")


@dataclass
class SkippedFile:
    path: str
    reason: str


@dataclass
class IngestionSummary:
    files_discovered: int = 0
    files_loaded: int = 0
    documents_produced: int = 0
    duplicate_count: int = 0
    skipped_files: list[SkippedFile] = field(default_factory=list)

    @property
    def files_skipped(self) -> int:
        return len(self.skipped_files)


def _infer_doc_type(filename: str) -> str:
    lowered = filename.lower()
    for keyword, doc_type in _DOC_TYPE_KEYWORDS:
        if keyword in lowered:
            return doc_type
    return "general"


def _infer_product_line(text: str) -> str:
    return "AquaPearl 500 Pro" if "aquapearl 500" in text.lower() else "general"


def normalize_whitespace(text: str) -> str:
    """Collapse redundant whitespace while preserving paragraph structure."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [_WHITESPACE_PATTERN.sub(" ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return _BLANK_LINES_PATTERN.sub("\n\n", text).strip()


def _extract_last_updated(text: str, file_mtime: float) -> str:
    match = _LAST_UPDATED_PATTERN.search(text)
    if match:
        return match.group(1).strip()
    return datetime.fromtimestamp(file_mtime).strftime("%Y-%m-%d")


def _content_hash(text: str) -> str:
    return hashlib.sha256(text.strip().lower().encode("utf-8")).hexdigest()


def _split_markdown_sections(text: str) -> list[tuple[str, str]]:
    """Split markdown/text content into (heading, body) sections."""
    lines = text.split("\n")
    sections: list[tuple[str, list[str], bool]] = []
    current_heading = "General"
    current_body: list[str] = []
    current_has_heading = False

    for line in lines:
        match = _HEADING_PATTERN.match(line)
        if match:
            if current_body:
                sections.append((current_heading, current_body, current_has_heading))
            current_heading = match.group(2).strip()
            current_body = []
            current_has_heading = True
        else:
            current_body.append(line)

    if current_body:
        sections.append((current_heading, current_body, current_has_heading))

    result = []
    for heading, body_lines, has_heading in sections:
        body_text = "\n".join(body_lines).strip()
        if not has_heading and _LAST_UPDATED_PATTERN.fullmatch(body_text.strip()):
            continue  # standalone "Last Updated:" preamble line, already captured in metadata
        if body_text:
            result.append((heading, f"{heading}\n{body_text}" if has_heading else body_text))
    return result or [("General", text)]


def _load_markdown_or_text(path: Path) -> list[Document]:
    raw = path.read_text(encoding="utf-8")
    normalized = normalize_whitespace(raw)
    last_updated = _extract_last_updated(raw, path.stat().st_mtime)
    doc_type = _infer_doc_type(path.name)
    product_line = _infer_product_line(normalized)

    documents = []
    for heading, section_text in _split_markdown_sections(normalized):
        if not section_text.strip():
            continue
        documents.append(
            Document(
                page_content=section_text,
                metadata={
                    "source": path.name,
                    "doc_type": doc_type,
                    "product_line": product_line,
                    "last_updated": last_updated,
                    "location": f"Section: {heading}",
                },
            )
        )
    return documents


def _load_docx(path: Path) -> list[Document]:
    document = docx.Document(str(path))
    last_updated = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")
    doc_type = _infer_doc_type(path.name)

    sections: list[tuple[str, list[str]]] = []
    current_heading = "General"
    current_body: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        is_heading = paragraph.style is not None and paragraph.style.name.lower().startswith(("heading", "title"))
        last_updated_match = _LAST_UPDATED_PATTERN.search(text)
        if last_updated_match:
            last_updated = last_updated_match.group(1).strip()
        if is_heading:
            if current_body:
                sections.append((current_heading, current_body))
            current_heading = text
            current_body = []
        else:
            current_body.append(text)
    if current_body:
        sections.append((current_heading, current_body))

    full_text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    product_line = _infer_product_line(full_text)

    documents = []
    for heading, body_lines in sections:
        body_text = normalize_whitespace("\n".join(body_lines))
        if not body_text:
            continue
        documents.append(
            Document(
                page_content=f"{heading}\n{body_text}",
                metadata={
                    "source": path.name,
                    "doc_type": doc_type,
                    "product_line": product_line,
                    "last_updated": last_updated,
                    "location": f"Section: {heading}",
                },
            )
        )
    return documents


def _load_pdf(path: Path) -> list[Document]:
    loader = PyPDFLoader(str(path))
    pages = loader.load()
    last_updated = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")
    doc_type = _infer_doc_type(path.name)

    documents = []
    for page in pages:
        normalized = normalize_whitespace(page.page_content)
        if not normalized:
            continue
        page_number = page.metadata.get("page", 0) + 1
        product_line = _infer_product_line(normalized)
        documents.append(
            Document(
                page_content=normalized,
                metadata={
                    "source": path.name,
                    "doc_type": doc_type,
                    "product_line": product_line,
                    "last_updated": last_updated,
                    "location": f"Page {page_number}",
                },
            )
        )
    return documents


def _load_csv(path: Path) -> list[Document]:
    last_updated = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")
    doc_type = _infer_doc_type(path.name)

    documents = []
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row_index, row in enumerate(reader, start=1):
            content = "\n".join(f"{key}: {value}" for key, value in row.items() if value)
            if not content.strip():
                continue
            sku = row.get("sku") or row.get("SKU") or row.get("model") or ""
            location = f"Row {row_index}" + (f" - SKU {sku}" if sku else "")
            product_line = _infer_product_line(content)
            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source": path.name,
                        "doc_type": doc_type,
                        "product_line": product_line,
                        "last_updated": last_updated,
                        "location": location,
                    },
                )
            )
    return documents


_LOADER_DISPATCH = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".md": _load_markdown_or_text,
    ".txt": _load_markdown_or_text,
    ".csv": _load_csv,
}


def load_knowledge_base(kb_dir: Path) -> tuple[list[Document], IngestionSummary]:
    """Load every supported document under kb_dir into normalized Documents."""
    summary = IngestionSummary()
    documents: list[Document] = []
    seen_hashes: set[str] = set()

    if not kb_dir.exists():
        logger.warning("Knowledge base directory does not exist: %s", kb_dir)
        return documents, summary

    paths = sorted(p for p in kb_dir.rglob("*") if p.is_file())
    summary.files_discovered = len(paths)

    for path in paths:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            summary.skipped_files.append(SkippedFile(str(path.name), f"Unsupported file type: {suffix}"))
            logger.info("Skipping unsupported file type: %s", path.name)
            continue

        try:
            file_documents = _LOADER_DISPATCH[suffix](path)
        except Exception as exc:  # noqa: BLE001 - must not abort the whole run
            summary.skipped_files.append(SkippedFile(str(path.name), f"Failed to load: {exc}"))
            logger.error("Failed to load %s: %s", path.name, exc, exc_info=True)
            continue

        if not file_documents:
            summary.skipped_files.append(SkippedFile(str(path.name), "No extractable content"))
            logger.info("No extractable content in %s", path.name)
            continue

        added_any = False
        for doc in file_documents:
            if not doc.page_content.strip():
                continue
            content_hash = _content_hash(doc.page_content)
            if content_hash in seen_hashes:
                summary.duplicate_count += 1
                logger.warning(
                    "Duplicate content detected, skipping: %s (%s)",
                    doc.metadata.get("source"),
                    doc.metadata.get("location"),
                )
                continue
            seen_hashes.add(content_hash)
            documents.append(doc)
            added_any = True

        if added_any:
            summary.files_loaded += 1

    summary.documents_produced = len(documents)
    logger.info(
        "Ingestion complete: %d files discovered, %d loaded, %d skipped, %d documents produced, %d duplicates",
        summary.files_discovered,
        summary.files_loaded,
        summary.files_skipped,
        summary.documents_produced,
        summary.duplicate_count,
    )
    return documents, summary
